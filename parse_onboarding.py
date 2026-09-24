"""Summarize client onboarding attempts and failures from a Cisco C9800 log."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict, deque
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import TextIO

from map_bssid import extract_bssids, update_inventory


MAC = r"[0-9a-fA-F]{4}(?:\.[0-9a-fA-F]{4}){2}"
TIMESTAMP_RE = re.compile(r"^(?P<timestamp>\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2})")
PRECISE_TIMESTAMP_RE = re.compile(
    r"^(?P<timestamp>\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2})"
    r"(?:\.(?P<fraction>\d{1,9}))?"
)
HKT = timezone(timedelta(hours=8), name="HKT")
AAA_LONG_LATENCY_MS = 10_000
AAA_LONG_LATENCY_REASON = "long latency >10s"
CLIENT_RE = re.compile(rf"\bMAC: (?P<client>{MAC})")
AAA_INIT_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+(?:MAB Authentication initiated\.|"
    r"L2 Authentication initiated\. method DOT1X\b)"
)
AAA_SUCCESS_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+(?:MAB Authentication success\.|"
    r"L2 Authentication Key Exchange Start\.)"
)
ASSOCIATION_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+"
    rf"(?:Re-)?Association received\. BSSID (?P<bssid>{MAC}), "
    rf"WLAN (?P<ssid>[^,]+), "
    rf"Slot [0-9]+ AP (?P<ap_mac>{MAC}), "
    r"(?P<ap_name>[^,]+), Site tag "
)
RUN_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+Client state transition: "
    r"S_CO_[A-Z0-9_]+ -> S_CO_RUN"
)
DELETE_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+Client delete initiated\. Reason: "
    r"(?P<reason>[A-Z0-9_]+)"
)
SEND_DELETE_RE = re.compile(
    rf"MAC: (?P<client>{MAC})\s+Send delete mobile\..*?"
    rf"BSSID: (?P<bssid>{MAC})"
)


@dataclass(frozen=True)
class AccessPoint:
    """Identity of an access point found in an association event."""

    name: str
    mac: str


@dataclass
class Counts:
    """Onboarding counters for one access point."""

    attempts: int = 0
    successes: int = 0
    failures: int = 0
    clients: set[str] = field(default_factory=set)
    failure_events: dict[str, int] = field(default_factory=dict)
    failure_clients: dict[str, set[str]] = field(default_factory=dict)
    client_failure_events: dict[str, dict[str, int]] = field(default_factory=dict)


@dataclass(frozen=True)
class Inventory:
    """Platform lookup indexes loaded from the AP inventory."""

    by_mac: dict[str, str]
    by_name: dict[str, str]
    devices_by_platform: dict[str, int] = field(default_factory=dict)
    bssids_by_ssid: dict[str, dict[str, AccessPoint]] = field(default_factory=dict)


@dataclass
class PlatformCounts(Counts):
    """Onboarding counters and AP count for one hardware platform."""

    access_points: int = 0
    association_events: int = 0
    client_ap_failure_events: dict[str, dict[str, int]] = field(default_factory=dict)


@dataclass(frozen=True)
class ClientLog:
    """One log record for a requested client."""

    line: str
    is_failure: bool
    failure_reason: str = ""
    onboarding_failure_explanation: str = ""
    aaa_failure_reason: str = ""


@dataclass(frozen=True)
class PendingFailure:
    """A deletion awaiting its delete-mobile BSSID record."""

    client: str
    reason: str
    access_point: AccessPoint | None
    client_log_index: int | None


@dataclass(frozen=True)
class AaaTransaction:
    """One MAB or DOT1X AAA transaction observed in the trace."""

    client: str
    start_time: str
    end_time: str
    latency_ms: float | None
    completed: bool
    reason: str = ""
    access_point: AccessPoint | None = None
    ssid: str = ""


@dataclass
class AaaClientCounts:
    """AAA completion and latency metrics for one client."""

    attempts: int = 0
    completed: int = 0
    incomplete: int = 0
    latencies_ms: list[float] = field(default_factory=list)
    incomplete_reasons: dict[str, int] = field(default_factory=dict)


@dataclass
class AaaAccessPointCounts:
    """AAA latency and incomplete metrics for one access point."""

    latencies_ms: list[float] = field(default_factory=list)
    incomplete: int = 0


@dataclass(frozen=True)
class Summary:
    """Parsed onboarding counters and capture time range."""

    counts: dict[AccessPoint, Counts]
    start_time: datetime | None
    stop_time: datetime | None
    association_events: dict[AccessPoint, int] = field(default_factory=dict)
    client_logs: dict[str, tuple[ClientLog, ...]] = field(default_factory=dict)
    aaa_transactions: tuple[AaaTransaction, ...] = ()


def normalize_mac(mac: str) -> str:
    """Return a MAC address without separators and in lowercase."""
    return "".join(character for character in mac.lower() if character.isalnum())


def parse_hkt_time(value: str) -> datetime:
    """Parse a command-line timestamp and attach the controller HKT timezone."""
    for time_format in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(value, time_format).replace(tzinfo=HKT)
        except ValueError:
            continue
    raise argparse.ArgumentTypeError(
        "expected YYYY-MM-DD HH:MM:SS or YYYY-MM-DDTHH:MM:SS"
    )


def parse_client_mac(value: str) -> str:
    """Validate and normalize a client MAC address to Cisco dotted notation."""
    normalized = normalize_mac(value)
    if len(normalized) != 12 or not all(
        character in "0123456789abcdef" for character in normalized
    ):
        raise argparse.ArgumentTypeError("expected a 12-digit hexadecimal MAC address")
    return ".".join(normalized[index : index + 4] for index in range(0, 12, 4))


def parse_client_macs(value: str) -> tuple[str, ...]:
    """Parse a comma-, semicolon-, or whitespace-separated client MAC list."""
    values = [item for item in re.split(r"[\s,;]+", value.strip()) if item]
    if not values:
        raise argparse.ArgumentTypeError("expected at least one client MAC address")
    return tuple(parse_client_mac(item) for item in values)


def trace_timestamp_ns(line: str) -> tuple[str, int] | None:
    """Return a display timestamp and integer nanoseconds for a trace line."""
    match = PRECISE_TIMESTAMP_RE.match(line)
    if not match:
        return None
    timestamp = match.group("timestamp")
    value = datetime.strptime(timestamp, "%Y/%m/%d %H:%M:%S")
    whole_seconds = (
        value.toordinal() * 86400
        + value.hour * 3600
        + value.minute * 60
        + value.second
    )
    fraction_value = match.group("fraction")
    fraction = (fraction_value or "").ljust(9, "0")
    display = timestamp + (f".{fraction_value}" if fraction_value else "")
    return display, whole_seconds * 1_000_000_000 + int(fraction or 0)


def load_inventory(stream: TextIO) -> Inventory:
    """Load AP platform mappings from a Catalyst Center inventory CSV."""
    reader = csv.DictReader(stream)
    required_columns = {"Device Name", "MAC Address", "Platform"}
    missing_columns = required_columns.difference(reader.fieldnames or ())
    if missing_columns:
        missing = ", ".join(sorted(missing_columns))
        raise ValueError(f"inventory is missing required columns: {missing}")

    by_mac: dict[str, str] = {}
    by_name: dict[str, str] = {}
    devices_by_platform: dict[str, int] = defaultdict(int)
    bssids_by_ssid: dict[str, dict[str, AccessPoint]] = defaultdict(dict)
    bssid_columns = [
        column for column in reader.fieldnames or () if column.endswith(" BSSID")
    ]
    for row in reader:
        platform = row["Platform"].strip() or "Unknown"
        devices_by_platform[platform] += 1
        access_point = AccessPoint(
            name=row["Device Name"].strip(),
            mac=row["MAC Address"].strip(),
        )
        if row["MAC Address"].strip():
            by_mac[normalize_mac(row["MAC Address"])] = platform
        if row["Device Name"].strip():
            by_name[row["Device Name"].strip()] = platform
        for column in bssid_columns:
            ssid = column.removesuffix(" BSSID")
            for bssid in row[column].split(";"):
                if bssid.strip():
                    bssids_by_ssid[ssid][normalize_mac(bssid)] = access_point
    return Inventory(
        by_mac=by_mac,
        by_name=by_name,
        devices_by_platform=dict(devices_by_platform),
        bssids_by_ssid=dict(bssids_by_ssid),
    )


def get_platform(access_point: AccessPoint, inventory: Inventory) -> str:
    """Find an AP platform by controller MAC, falling back to AP name."""
    return inventory.by_mac.get(
        normalize_mac(access_point.mac),
        inventory.by_name.get(access_point.name, "Not in inventory"),
    )


def summarize(
    stream: TextIO,
    ssid: str = "Corp@CLP",
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    client_macs: tuple[str, ...] | str | None = None,
    inventory: Inventory | None = None,
) -> Summary:
    """Aggregate onboarding attempts and terminal failures from log lines.

    An association or reassociation starts an attempt. The attempt succeeds
    when the client state reaches ``S_CO_RUN`` and fails when the client is
    deleted first. Delete reasons containing ``SUCCESS`` are not failures.

    Args:
        stream: Text stream containing C9800 wireless profile log lines.
        ssid: Exact WLAN name to include.
        start_time: Inclusive lower HKT boundary for log records.
        end_time: Inclusive upper HKT boundary for log records.
        client_macs: Optional normalized client MACs whose records are collected.
        inventory: Optional AP inventory containing SSID BSSID mappings.

    Returns:
        Per-access-point counters and the log capture time range in HKT.
    """
    counts: dict[AccessPoint, Counts] = defaultdict(Counts)
    association_events: dict[AccessPoint, int] = defaultdict(int)
    active_attempts: dict[str, AccessPoint] = {}
    pending_failures: dict[str, PendingFailure] = {}
    pending_aaa: dict[
        str,
        deque[
            tuple[
                str,
                int,
                str,
                AccessPoint | None,
                str,
                int | None,
                int | None,
            ]
        ],
    ] = defaultdict(deque)
    aaa_transactions: list[AaaTransaction] = []
    association_context: dict[str, tuple[AccessPoint, str]] = {}
    first_timestamp: str | None = None
    last_timestamp: str | None = None
    requested_clients = (
        (client_macs,) if isinstance(client_macs, str) else (client_macs or ())
    )
    requested_clients = tuple(dict.fromkeys(client.lower() for client in requested_clients))
    requested_client_set = set(requested_clients)
    client_logs: dict[str, list[ClientLog]] = {
        client: [] for client in requested_clients
    }
    start_value = start_time.strftime("%Y/%m/%d %H:%M:%S") if start_time else None
    end_value = end_time.strftime("%Y/%m/%d %H:%M:%S") if end_time else None
    ssid_bssids = inventory.bssids_by_ssid.get(ssid, {}) if inventory else {}

    def resolve_failure(client: str, bssid: str | None = None) -> None:
        pending = pending_failures.pop(client, None)
        if pending is None:
            return

        access_point = pending.access_point
        counted = access_point is not None
        if access_point is None:
            mapped_access_point = (
                ssid_bssids.get(normalize_mac(bssid)) if bssid else None
            )
            if mapped_access_point:
                explanation = (
                    f"Not counted: BSSID {bssid.lower()} maps to {ssid} on AP "
                    f"{mapped_access_point.name}, but no onboarding attempt is active"
                )
            elif bssid and ssid_bssids:
                explanation = (
                    f"Not counted: BSSID {bssid.lower()} is not mapped to {ssid} "
                    "and no onboarding attempt is active"
                )
            else:
                explanation = f"Not counted: no active {ssid} onboarding attempt"
        elif bssid and ssid_bssids:
            mapped_access_point = ssid_bssids.get(normalize_mac(bssid))
            counted = bool(
                mapped_access_point
                and normalize_mac(mapped_access_point.mac)
                == normalize_mac(access_point.mac)
            )
            if counted:
                explanation = (
                    f"Counted: BSSID {bssid.lower()} confirms the active {ssid} "
                    f"onboarding attempt on AP {access_point.name}"
                )
            elif mapped_access_point:
                explanation = (
                    f"Not counted: BSSID {bssid.lower()} maps to AP "
                    f"{mapped_access_point.name}, not the active attempt on "
                    f"{access_point.name}"
                )
            else:
                explanation = (
                    f"Not counted: BSSID {bssid.lower()} is not mapped to {ssid}"
                )
        else:
            explanation = (
                f"Counted: active {ssid} onboarding attempt on AP "
                f"{access_point.name}; no mapped deletion BSSID was available"
            )

        if counted and access_point:
            if active_attempts.get(client) == access_point:
                active_attempts.pop(client)
            result = counts[access_point]
            result.failures += 1
            result.failure_events[pending.reason] = (
                result.failure_events.get(pending.reason, 0) + 1
            )
            result.failure_clients.setdefault(pending.reason, set()).add(client)
            client_reasons = result.client_failure_events.setdefault(client, {})
            client_reasons[pending.reason] = client_reasons.get(pending.reason, 0) + 1

        if pending.client_log_index is not None:
            entry = client_logs[pending.client][pending.client_log_index]
            client_logs[pending.client][pending.client_log_index] = ClientLog(
                entry.line,
                entry.is_failure,
                entry.failure_reason,
                explanation,
                entry.aaa_failure_reason,
            )

    def resolve_aaa(
        client: str,
        end_timestamp: str,
        end_ns: int,
        completed: bool,
        reason: str = "",
        end_client_log_index: int | None = None,
    ) -> None:
        client_pending = pending_aaa.get(client)
        if not client_pending:
            return
        (
            start_timestamp,
            start_ns,
            pending_reason,
            access_point,
            aaa_ssid,
            start_client_log_index,
            pending_stop_log_index,
        ) = client_pending.popleft()
        if not client_pending:
            pending_aaa.pop(client, None)
        latency_ms = (end_ns - start_ns) / 1_000_000 if completed else None
        is_long_latency = bool(
            completed
            and latency_ms is not None
            and latency_ms > AAA_LONG_LATENCY_MS
        )
        transaction_completed = completed and not is_long_latency
        failure_reason = (
            AAA_LONG_LATENCY_REASON if is_long_latency else reason or pending_reason
        )
        aaa_transactions.append(
            AaaTransaction(
                client=client,
                start_time=start_timestamp,
                end_time=end_timestamp,
                latency_ms=latency_ms,
                completed=transaction_completed,
                reason=failure_reason,
                access_point=access_point,
                ssid=aaa_ssid,
            )
        )
        if not transaction_completed and failure_reason:
            stop_client_log_index = (
                pending_stop_log_index
                if pending_stop_log_index is not None
                else end_client_log_index
            )
            for log_index in {start_client_log_index, stop_client_log_index}:
                if log_index is None:
                    continue
                entry = client_logs[client][log_index]
                client_logs[client][log_index] = ClientLog(
                    line=entry.line,
                    is_failure=True,
                    failure_reason=entry.failure_reason,
                    onboarding_failure_explanation=entry.onboarding_failure_explanation,
                    aaa_failure_reason=f"AAA failure: {failure_reason}",
                )

    for line in stream:
        precise_timestamp = trace_timestamp_ns(line)
        timestamp = TIMESTAMP_RE.match(line)
        if timestamp:
            line_timestamp = timestamp.group("timestamp")
            if start_value and line_timestamp < start_value:
                continue
            if end_value and line_timestamp > end_value:
                break
            last_timestamp = line_timestamp
            if first_timestamp is None:
                first_timestamp = last_timestamp
        elif start_value or end_value:
            continue

        client_log_index: int | None = None
        if requested_client_set:
            client_match = CLIENT_RE.search(line)
            matched_client = client_match.group("client").lower() if client_match else None
            if matched_client in requested_client_set:
                deletion = DELETE_RE.search(line)
                is_failure = bool(
                    deletion and "SUCCESS" not in deletion.group("reason")
                )
                failure_reason = ""
                onboarding_failure_explanation = ""
                if is_failure and deletion:
                    failure_reason = deletion.group("reason").removeprefix(
                        "CO_CLIENT_DELETE_REASON_"
                    )
                    onboarding_failure_explanation = "Pending BSSID validation"
                client_logs[matched_client].append(
                    ClientLog(
                        line.rstrip("\r\n"),
                        is_failure,
                        failure_reason,
                        onboarding_failure_explanation,
                    )
                )
                client_log_index = len(client_logs[matched_client]) - 1

        aaa_init = AAA_INIT_RE.search(line)
        if aaa_init and precise_timestamp:
            client = aaa_init.group("client").lower()
            access_point, aaa_ssid = association_context.get(client, (None, ""))
            pending_aaa[client].append(
                (
                    *precise_timestamp,
                    "",
                    access_point,
                    aaa_ssid,
                    client_log_index if client in requested_client_set else None,
                    None,
                )
            )
            continue

        aaa_success = AAA_SUCCESS_RE.search(line)
        if aaa_success and precise_timestamp:
            resolve_aaa(
                aaa_success.group("client").lower(),
                precise_timestamp[0],
                precise_timestamp[1],
                True,
                end_client_log_index=(
                    client_log_index
                    if aaa_success.group("client").lower() in requested_client_set
                    else None
                ),
            )
            continue

        send_delete = SEND_DELETE_RE.search(line)
        if send_delete:
            resolve_failure(
                send_delete.group("client").lower(),
                send_delete.group("bssid"),
            )
            continue

        association = ASSOCIATION_RE.search(line)
        if association:
            client = association.group("client").lower()
            while client in pending_aaa and precise_timestamp:
                pending_reason = pending_aaa[client][0][2]
                resolve_aaa(
                    client,
                    precise_timestamp[0],
                    precise_timestamp[1],
                    False,
                    pending_reason or "Superseded by a new association",
                    client_log_index if client in requested_client_set else None,
                )
            resolve_failure(client)
            access_point = AccessPoint(
                name=association.group("ap_name").strip(),
                mac=association.group("ap_mac").lower(),
            )
            association_context[client] = (
                access_point,
                association.group("ssid").strip(),
            )
            association_events[access_point] += 1
            if association.group("ssid").strip() != ssid:
                active_attempts.pop(client, None)
                continue
            if ssid_bssids:
                mapped_access_point = ssid_bssids.get(
                    normalize_mac(association.group("bssid"))
                )
                if not mapped_access_point or normalize_mac(
                    mapped_access_point.mac
                ) != normalize_mac(access_point.mac):
                    active_attempts.pop(client, None)
                    continue
            counts[access_point].attempts += 1
            counts[access_point].clients.add(client)
            active_attempts[client] = access_point
            continue

        run = RUN_RE.search(line)
        if run:
            client = run.group("client").lower()
            resolve_failure(client)
            access_point = active_attempts.pop(client, None)
            if access_point:
                counts[access_point].successes += 1
            continue

        deletion = DELETE_RE.search(line)
        if deletion:
            client = deletion.group("client").lower()
            if client in pending_aaa:
                (
                    start_timestamp,
                    start_ns,
                    _,
                    access_point,
                    aaa_ssid,
                    start_client_log_index,
                    _,
                ) = pending_aaa[client][0]
                pending_aaa[client][0] = (
                    start_timestamp,
                    start_ns,
                    deletion.group("reason").removeprefix(
                        "CO_CLIENT_DELETE_REASON_"
                    ),
                    access_point,
                    aaa_ssid,
                    start_client_log_index,
                    client_log_index if client in requested_client_set else None,
                )
            resolve_failure(client)
            access_point = active_attempts.get(client)
            if "SUCCESS" not in deletion.group("reason"):
                pending_failures[client] = PendingFailure(
                    client=client,
                    reason=deletion.group("reason"),
                    access_point=access_point,
                    client_log_index=client_log_index,
                )

    for client in tuple(pending_failures):
        resolve_failure(client)
    capture_end = last_timestamp or "End of capture"
    for client in tuple(pending_aaa):
        while client in pending_aaa:
            _, start_ns, reason, _, _, _, _ = pending_aaa[client][0]
            resolve_aaa(
                client,
                capture_end,
                start_ns,
                False,
                reason or "No AAA completion record before end of capture",
            )

    def parse_time(value: str | None) -> datetime | None:
        if value is None:
            return None
        return datetime.strptime(value, "%Y/%m/%d %H:%M:%S").replace(tzinfo=HKT)

    return Summary(
        counts=dict(counts),
        association_events=dict(association_events),
        start_time=parse_time(first_timestamp),
        stop_time=parse_time(last_timestamp),
        client_logs={
            client: tuple(client_logs[client]) for client in requested_clients
        },
        aaa_transactions=tuple(aaa_transactions),
    )


def aggregate_aaa_by_client(
    transactions: tuple[AaaTransaction, ...],
) -> dict[str, AaaClientCounts]:
    """Aggregate AAA completion and latency metrics by client MAC."""
    clients: dict[str, AaaClientCounts] = defaultdict(AaaClientCounts)
    for transaction in transactions:
        result = clients[transaction.client]
        result.attempts += 1
        if transaction.completed and transaction.latency_ms is not None:
            result.completed += 1
            result.latencies_ms.append(transaction.latency_ms)
        else:
            result.incomplete += 1
            result.incomplete_reasons[transaction.reason] = (
                result.incomplete_reasons.get(transaction.reason, 0) + 1
            )
    return dict(clients)


def aggregate_aaa_by_client_and_ssid(
    transactions: tuple[AaaTransaction, ...],
) -> dict[tuple[str, str], AaaClientCounts]:
    """Aggregate AAA completion and latency metrics by client and SSID."""
    clients: dict[tuple[str, str], AaaClientCounts] = defaultdict(AaaClientCounts)
    for transaction in transactions:
        result = clients[(transaction.client, transaction.ssid)]
        result.attempts += 1
        if transaction.completed and transaction.latency_ms is not None:
            result.completed += 1
            result.latencies_ms.append(transaction.latency_ms)
        else:
            result.incomplete += 1
            result.incomplete_reasons[transaction.reason] = (
                result.incomplete_reasons.get(transaction.reason, 0) + 1
            )
    return dict(clients)


def aggregate_aaa_by_access_point(
    transactions: tuple[AaaTransaction, ...], ssid: str
) -> dict[AccessPoint, AaaAccessPointCounts]:
    """Aggregate selected-SSID AAA metrics by access point."""
    access_points: dict[AccessPoint, AaaAccessPointCounts] = defaultdict(
        AaaAccessPointCounts
    )
    for transaction in transactions:
        if transaction.access_point is None or transaction.ssid != ssid:
            continue
        result = access_points[transaction.access_point]
        if transaction.completed and transaction.latency_ms is not None:
            result.latencies_ms.append(transaction.latency_ms)
        else:
            result.incomplete += 1
    return dict(access_points)


def aggregate_aaa_by_platform(
    transactions: tuple[AaaTransaction, ...],
    ssid: str,
    inventory: Inventory,
) -> dict[str, AaaAccessPointCounts]:
    """Aggregate selected-SSID AAA metrics by AP hardware platform."""
    platforms: dict[str, AaaAccessPointCounts] = defaultdict(AaaAccessPointCounts)
    for transaction in transactions:
        if transaction.access_point is None or transaction.ssid != ssid:
            continue
        result = platforms[get_platform(transaction.access_point, inventory)]
        if transaction.completed and transaction.latency_ms is not None:
            result.latencies_ms.append(transaction.latency_ms)
        else:
            result.incomplete += 1
    return dict(platforms)


def aggregate_aaa_access_points_by_client_and_ssid(
    transactions: tuple[AaaTransaction, ...],
) -> dict[tuple[str, str], dict[str, int]]:
    """Count AAA transactions by client, SSID, and access point name."""
    clients: dict[tuple[str, str], dict[str, int]] = defaultdict(
        lambda: defaultdict(int)
    )
    for transaction in transactions:
        access_point_name = (
            transaction.access_point.name if transaction.access_point else "Unknown"
        )
        clients[(transaction.client, transaction.ssid)][access_point_name] += 1
    return {key: dict(access_points) for key, access_points in clients.items()}


def format_aaa_summary(transactions: tuple[AaaTransaction, ...]) -> str:
    """Format overall AAA completion and latency metrics."""
    completed = [
        transaction.latency_ms
        for transaction in transactions
        if transaction.completed and transaction.latency_ms is not None
    ]
    incomplete_reasons: dict[str, int] = defaultdict(int)
    for transaction in transactions:
        if not transaction.completed:
            incomplete_reasons[transaction.reason] += 1
    incomplete = len(transactions) - len(completed)
    completion_rate = len(completed) / len(transactions) * 100 if transactions else 0
    lines = [
        f"Transactions: {len(transactions)}",
        f"Completed: {len(completed)}",
        f"Incomplete: {incomplete}",
        f"Completion: {completion_rate:.2f}%",
        (
            f"Average latency: {sum(completed) / len(completed):.3f} ms"
            if completed
            else "Average latency: N/A"
        ),
        f"Maximum latency: {max(completed):.3f} ms" if completed else "Maximum latency: N/A",
    ]
    if incomplete_reasons:
        lines.append("Incomplete reasons:")
        lines.extend(
            f"  {reason}: {count}"
            for reason, count in sorted(
                incomplete_reasons.items(), key=lambda item: (-item[1], item[0])
            )
        )
    return "\n".join(lines)


def format_aaa_clients(
    transactions: tuple[AaaTransaction, ...], clients: tuple[str, ...]
) -> str:
    """Format AAA metrics for requested client MAC addresses."""
    headers = (
        "Client MAC",
        "Transactions",
        "Completed",
        "Incomplete",
        "Average ms",
        "Maximum ms",
        "Incomplete Reasons",
    )
    counts = aggregate_aaa_by_client(transactions)
    rows = []
    for client in clients:
        result = counts.get(client, AaaClientCounts())
        reasons = ", ".join(
            f"{reason} ({count})"
            for reason, count in sorted(
                result.incomplete_reasons.items(),
                key=lambda item: (-item[1], item[0]),
            )
        )
        rows.append(
            (
                client,
                str(result.attempts),
                str(result.completed),
                str(result.incomplete),
                (
                    f"{sum(result.latencies_ms) / len(result.latencies_ms):.3f}"
                    if result.latencies_ms
                    else "N/A"
                ),
                f"{max(result.latencies_ms):.3f}" if result.latencies_ms else "N/A",
                reasons,
            )
        )
    widths = [
        max([len(headers[index]), *(len(row[index]) for row in rows)])
        for index in range(len(headers))
    ]

    def render(row: tuple[str, ...]) -> str:
        return " | ".join(
            value.ljust(widths[index]) if index in (0, 6) else value.rjust(widths[index])
            for index, value in enumerate(row)
        )

    separator = "-+-".join("-" * width for width in widths)
    return "\n".join((render(headers), separator, *(render(row) for row in rows)))


def format_ap_table(counts: dict[AccessPoint, Counts], inventory: Inventory) -> str:
    """Format access-point counters as an aligned text table."""
    headers = (
        "AP Name",
        "AP MAC",
        "Platform",
        "Client MACs",
        "Attempts",
        "Successes",
        "Failures",
        "Success %",
    )
    rows = [
        (
            access_point.name,
            access_point.mac,
            get_platform(access_point, inventory),
            str(len(result.clients)),
            str(result.attempts),
            str(result.successes),
            str(result.failures),
            f"{result.successes / result.attempts * 100:.2f}",
        )
        for access_point, result in sorted(
            counts.items(), key=lambda item: (-item[1].attempts, item[0].name)
        )
    ]
    widths = [
        max([len(headers[index]), *(len(row[index]) for row in rows)])
        for index in range(len(headers))
    ]

    def render(row: tuple[str, ...]) -> str:
        return " | ".join(
            value.ljust(widths[index]) if index < 3 else value.rjust(widths[index])
            for index, value in enumerate(row)
        )

    separator = "-+-".join("-" * width for width in widths)
    return "\n".join((render(headers), separator, *(render(row) for row in rows)))


def aggregate_by_platform(
    counts: dict[AccessPoint, Counts],
    association_events: dict[AccessPoint, int],
    inventory: Inventory,
) -> dict[str, PlatformCounts]:
    """Combine per-AP onboarding counters by hardware platform."""
    platforms: dict[str, PlatformCounts] = defaultdict(PlatformCounts)
    observed_access_points: dict[str, set[AccessPoint]] = defaultdict(set)
    for access_point, result in counts.items():
        platform_counts = platforms[get_platform(access_point, inventory)]
        platform_counts.attempts += result.attempts
        platform_counts.successes += result.successes
        platform_counts.failures += result.failures
        platform_counts.clients.update(result.clients)
        for reason, failures in result.failure_events.items():
            platform_counts.failure_events[reason] = (
                platform_counts.failure_events.get(reason, 0) + failures
            )
            platform_counts.failure_clients.setdefault(reason, set()).update(
                result.failure_clients[reason]
            )
        for client, reasons in result.client_failure_events.items():
            platform_client_reasons = platform_counts.client_failure_events.setdefault(
                client, {}
            )
            for reason, failures in reasons.items():
                platform_client_reasons[reason] = (
                    platform_client_reasons.get(reason, 0) + failures
                )
            client_access_points = platform_counts.client_ap_failure_events.setdefault(
                client, {}
            )
            client_access_points[access_point.name] = (
                client_access_points.get(access_point.name, 0) + sum(reasons.values())
            )
    for access_point, events in association_events.items():
        platform = get_platform(access_point, inventory)
        if platform in platforms:
            platforms[platform].association_events += events
            observed_access_points[platform].add(access_point)
    for platform, platform_counts in platforms.items():
        platform_counts.access_points = inventory.devices_by_platform.get(
            platform, len(observed_access_points[platform])
        )
    return dict(platforms)


def format_platform_table(counts: dict[str, PlatformCounts]) -> str:
    """Format aggregate onboarding counters by hardware platform."""
    headers = (
        "Platform",
        "APs",
        "Client MACs",
        "Attempts",
        "Successes",
        "Failures",
        "Success %",
        "Failure %",
        "Unknown %",
        "Association Events",
        "Association Rate all SSID (per AP)",
    )
    rows = [
        (
            platform,
            str(result.access_points),
            str(len(result.clients)),
            str(result.attempts),
            str(result.successes),
            str(result.failures),
            f"{result.successes / result.attempts * 100:.2f}",
            f"{result.failures / result.attempts * 100:.2f}",
            f"{(result.attempts - result.successes - result.failures) / result.attempts * 100:.2f}",
            str(result.association_events),
            f"{result.association_events / result.access_points:.2f}",
        )
        for platform, result in sorted(counts.items())
    ]
    widths = [
        max(len(headers[index]), *(len(row[index]) for row in rows))
        for index in range(len(headers))
    ]

    def render(row: tuple[str, ...]) -> str:
        return " | ".join(
            value.ljust(widths[index]) if index == 0 else value.rjust(widths[index])
            for index, value in enumerate(row)
        )

    separator = "-+-".join("-" * width for width in widths)
    return "\n".join((render(headers), separator, *(render(row) for row in rows)))


def failure_rows(
    counts: dict[str, PlatformCounts], show_macs: bool
) -> list[tuple[str, ...]]:
    """Build failure-reason rows with optional client MAC address lists."""
    rows = [
        (
            platform,
            reason.removeprefix("CO_CLIENT_DELETE_REASON_"),
            str(result.failure_events[reason]),
            str(len(result.failure_clients[reason])),
            *(
                (", ".join(sorted(result.failure_clients[reason])),)
                if show_macs
                else ()
            ),
        )
        for platform, result in sorted(counts.items())
        for reason in sorted(
            result.failure_events,
            key=lambda value: (-result.failure_events[value], value),
        )
    ]
    return rows


def format_failure_table(
    counts: dict[str, PlatformCounts], show_macs: bool = False
) -> str:
    """Format failure reasons and affected client MACs by platform."""
    headers = ("Platform", "Failure Reason", "Failures", "Client MACs")
    if show_macs:
        headers += ("MAC Addresses",)
    rows = failure_rows(counts, show_macs)
    widths = [
        max([len(headers[index]), *(len(row[index]) for row in rows)])
        for index in range(len(headers))
    ]

    def render(row: tuple[str, ...]) -> str:
        return " | ".join(
            value.rjust(widths[index]) if index in (2, 3) else value.ljust(widths[index])
            for index, value in enumerate(row)
        )

    separator = "-+-".join("-" * width for width in widths)
    return "\n".join((render(headers), separator, *(render(row) for row in rows)))


def top_error_client_rows(
    counts: dict[str, PlatformCounts], limit: int
) -> list[tuple[str, str, str, str]]:
    """Build ranked client failure rows for each platform."""
    rows: list[tuple[str, str, str, str]] = []
    for platform, result in sorted(counts.items()):
        ranked_clients = sorted(
            result.client_failure_events.items(),
            key=lambda item: (-sum(item[1].values()), item[0]),
        )[:limit]
        for client, reasons in ranked_clients:
            reason_summary = ", ".join(
                f"{reason.removeprefix('CO_CLIENT_DELETE_REASON_')} ({failures})"
                for reason, failures in sorted(
                    reasons.items(), key=lambda item: (-item[1], item[0])
                )
            )
            rows.append(
                (platform, client, str(sum(reasons.values())), reason_summary)
            )
    return rows


def top_error_client_macs(
    counts: dict[str, PlatformCounts], limit: int
) -> tuple[str, ...]:
    """Return the first ranked failing client MACs from each platform."""
    return tuple(
        dict.fromkeys(row[1] for row in top_error_client_rows(counts, limit))
    )


def format_top_error_clients(
    counts: dict[str, PlatformCounts], limit: int
) -> str:
    """Format the clients with the most failure events per platform."""
    headers = ("Platform", "Client MAC", "Failures", "Failure Reasons")
    rows = top_error_client_rows(counts, limit)
    widths = [
        max([len(headers[index]), *(len(row[index]) for row in rows)])
        for index in range(len(headers))
    ]

    def render(row: tuple[str, ...]) -> str:
        return " | ".join(
            value.rjust(widths[index]) if index == 2 else value.ljust(widths[index])
            for index, value in enumerate(row)
        )

    separator = "-+-".join("-" * width for width in widths)
    return "\n".join((render(headers), separator, *(render(row) for row in rows)))


def save_xlsx(
    output_path: Path,
    summary: Summary,
    inventory: Inventory,
    platform_counts: dict[str, PlatformCounts],
    ssid: str,
    show_macs: bool,
    top_error_clients: int | None,
    client_macs: tuple[str, ...] | None,
) -> None:
    """Save the onboarding report to a formatted XLSX workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    metadata = workbook.active
    if metadata is None:
        raise RuntimeError("failed to create workbook")
    metadata.title = "Summary"
    time_format = "%Y-%m-%d %H:%M:%S %Z"
    metadata_rows = [
        ("SSID", ssid),
        (
            "Start time",
            summary.start_time.strftime(time_format) if summary.start_time else "Unknown",
        ),
        (
            "Stop time",
            summary.stop_time.strftime(time_format) if summary.stop_time else "Unknown",
        ),
    ]
    if client_macs:
        metadata_rows.append(("Client log MACs", ", ".join(client_macs)))
    for row in metadata_rows:
        metadata.append(row)

    per_ap = workbook.create_sheet("Per AP")
    aaa_by_access_point = aggregate_aaa_by_access_point(
        summary.aaa_transactions, ssid
    )
    per_ap.append(
        (
            "AP Name",
            "AP MAC",
            "Platform",
            "Client MACs",
            "Attempts",
            "Successes",
            "Failures",
            "Success %",
            "AAA Average Latency (ms)",
            "AAA Incomplete",
        )
    )
    for access_point, result in sorted(
        summary.counts.items(), key=lambda item: (-item[1].attempts, item[0].name)
    ):
        aaa_result = aaa_by_access_point.get(access_point, AaaAccessPointCounts())
        per_ap.append(
            (
                access_point.name,
                access_point.mac,
                get_platform(access_point, inventory),
                len(result.clients),
                result.attempts,
                result.successes,
                result.failures,
                result.successes / result.attempts,
                (
                    sum(aaa_result.latencies_ms) / len(aaa_result.latencies_ms)
                    if aaa_result.latencies_ms
                    else None
                ),
                aaa_result.incomplete,
            )
        )

    by_platform = workbook.create_sheet("By Platform")
    aaa_by_platform = aggregate_aaa_by_platform(
        summary.aaa_transactions, ssid, inventory
    )
    by_platform.append(
        (
            "Platform",
            "APs",
            "Client MACs",
            "Attempts",
            "Successes",
            "Failures",
            "Success %",
            "Failure %",
            "Unknown %",
            "Association Events",
            "Association Rate all SSID (per AP)",
            "AAA Average Latency (ms)",
            "AAA Incomplete",
            "AAA Average Incomplete (per AP)",
        )
    )
    for platform, result in sorted(platform_counts.items()):
        aaa_result = aaa_by_platform.get(platform, AaaAccessPointCounts())
        by_platform.append(
            (
                platform,
                result.access_points,
                len(result.clients),
                result.attempts,
                result.successes,
                result.failures,
                result.successes / result.attempts,
                result.failures / result.attempts,
                (result.attempts - result.successes - result.failures)
                / result.attempts,
                result.association_events,
                result.association_events / result.access_points,
                (
                    sum(aaa_result.latencies_ms) / len(aaa_result.latencies_ms)
                    if aaa_result.latencies_ms
                    else None
                ),
                aaa_result.incomplete,
                (
                    aaa_result.incomplete / result.access_points
                    if result.access_points
                    else None
                ),
            )
        )

    aaa_by_client = workbook.create_sheet("AAA by Client")
    aaa_by_client.append(
        (
            "SSID",
            "Client MAC",
            "Transactions",
            "Completed",
            "Incomplete",
            "Completion %",
            "Average Latency (ms)",
            "Minimum Latency (ms)",
            "Maximum Latency (ms)",
            "Incomplete Reasons",
            "AP Transactions",
        )
    )
    aaa_access_points_by_client_and_ssid = (
        aggregate_aaa_access_points_by_client_and_ssid(
            summary.aaa_transactions
        )
    )
    for (client, client_ssid), result in sorted(
        aggregate_aaa_by_client_and_ssid(summary.aaa_transactions).items()
    ):
        reasons = ", ".join(
            f"{reason} ({count})"
            for reason, count in sorted(
                result.incomplete_reasons.items(),
                key=lambda item: (-item[1], item[0]),
            )
        )
        access_point_summary = ", ".join(
            f"{access_point} ({transactions})"
            for access_point, transactions in sorted(
                aaa_access_points_by_client_and_ssid.get(
                    (client, client_ssid), {}
                ).items(),
                key=lambda item: (-item[1], item[0]),
            )
        )
        aaa_by_client.append(
            (
                client_ssid,
                client,
                result.attempts,
                result.completed,
                result.incomplete,
                result.completed / result.attempts,
                (
                    sum(result.latencies_ms) / len(result.latencies_ms)
                    if result.latencies_ms
                    else None
                ),
                min(result.latencies_ms) if result.latencies_ms else None,
                max(result.latencies_ms) if result.latencies_ms else None,
                reasons,
                access_point_summary,
            )
        )

    incomplete_aaa = [
        transaction
        for transaction in summary.aaa_transactions
        if not transaction.completed
    ]
    if incomplete_aaa:
        aaa_incomplete = workbook.create_sheet("AAA Incomplete")
        aaa_incomplete.append(
            (
                "SSID",
                "Client MAC",
                "AP Name",
                "Platform",
                "Start Time",
                "End Time",
                "Latency (ms)",
                "Reason",
            )
        )
        for transaction in incomplete_aaa:
            aaa_incomplete.append(
                (
                    transaction.ssid,
                    transaction.client,
                    transaction.access_point.name if transaction.access_point else "Unknown",
                    (
                        get_platform(transaction.access_point, inventory)
                        if transaction.access_point
                        else "Unknown"
                    ),
                    transaction.start_time,
                    transaction.end_time,
                    transaction.latency_ms,
                    transaction.reason,
                )
            )

    failures = workbook.create_sheet("Failure Reasons")
    failure_headers = (
        "Platform",
        "Failure Reason",
        "Failures",
        "Unique Clients for Reason",
        "Unique Failed Clients for Platform",
    )
    if show_macs:
        failure_headers += ("MAC Addresses",)
    failures.append(failure_headers)
    for row in failure_rows(platform_counts, show_macs):
        failures.append(
            (
                *row[:2],
                int(row[2]),
                int(row[3]),
                len(platform_counts[row[0]].client_failure_events),
                *row[4:],
            )
        )

    if top_error_clients is not None:
        top_clients = workbook.create_sheet("Top Error Clients")
        top_clients.append(
            ("Platform", "Client MAC", "Failures", "Failure Reasons", "AP Failures")
        )
        for row in top_error_client_rows(platform_counts, top_error_clients):
            access_points = platform_counts[row[0]].client_ap_failure_events[row[1]]
            access_point_summary = ", ".join(
                f"{access_point} ({failures})"
                for access_point, failures in sorted(
                    access_points.items(), key=lambda item: (-item[1], item[0])
                )
            )
            top_clients.append((*row[:2], int(row[2]), row[3], access_point_summary))

    for client_mac in client_macs or ():
        client_log_sheet = workbook.create_sheet(f"Client {client_mac}")
        client_log_sheet.append(
            (
                "Failure Event",
                "Failure Reason",
                "Counted Onboarding Failure",
                "Log",
            )
        )
        for entry in summary.client_logs.get(client_mac, ()):
            failure_reason = "; ".join(
                reason
                for reason in (entry.failure_reason, entry.aaa_failure_reason)
                if reason
            )
            client_log_sheet.append(
                (
                    "Yes" if entry.is_failure else "No",
                    failure_reason,
                    entry.onboarding_failure_explanation,
                    entry.line,
                )
            )
            if entry.is_failure:
                for cell in client_log_sheet[client_log_sheet.max_row]:
                    cell.font = Font(bold=True, color="FFFF0000")

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column_cells in worksheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column_cells) + 2,
                80,
            )
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for cell in per_ap["H"][1:]:
        cell.number_format = "0.00%"
    for cell in per_ap["I"][1:]:
        cell.number_format = "0.000"
    for cell in by_platform["K"][1:]:
        cell.number_format = "0.00"
    for cell in by_platform["L"][1:]:
        cell.number_format = "0.000"
    for cell in by_platform["N"][1:]:
        cell.number_format = "0.00"
    if incomplete_aaa:
        for cell in aaa_incomplete["G"][1:]:
            cell.number_format = "0.000"
    for cell in aaa_by_client["F"][1:]:
        cell.number_format = "0.00%"
    for column in ("G", "H", "I"):
        for cell in aaa_by_client[column][1:]:
            cell.number_format = "0.000"
    for column in ("G", "H", "I"):
        for cell in by_platform[column][1:]:
            cell.number_format = "0.00%"

    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Summarize C9800 client onboarding attempts and failures by AP."
    )
    parser.add_argument("log_file", type=Path, help="C9800 wireless profile log")
    parser.add_argument(
        "--ssid",
        default="Corp@CLP",
        help="exact WLAN name to include (default: %(default)s)",
    )
    parser.add_argument(
        "--inventory",
        type=Path,
        default=Path("clp_inventory_2f.csv"),
        help="AP inventory CSV (default: %(default)s)",
    )
    parser.add_argument(
        "--start-time",
        type=parse_hkt_time,
        metavar="TIMESTAMP",
        help="inclusive HKT start time (YYYY-MM-DD HH:MM:SS)",
    )
    parser.add_argument(
        "--end-time",
        type=parse_hkt_time,
        metavar="TIMESTAMP",
        help="inclusive HKT end time (YYYY-MM-DD HH:MM:SS)",
    )
    parser.add_argument(
        "--show-macs",
        action="store_true",
        help="show affected client MAC address lists (hidden by default)",
    )
    parser.add_argument(
        "--top-error-clients",
        type=int,
        metavar="N",
        help="show the top N failing client MACs per platform",
    )
    parser.add_argument(
        "--top-error-client-logs",
        type=int,
        metavar="N",
        help=(
            "write client-log worksheets for the first N clients per platform "
            "in Top Error Clients"
        ),
    )
    parser.add_argument(
        "--client-logs",
        action="append",
        type=parse_client_macs,
        metavar="MAC[,MAC...]",
        help=(
            "write each client's log records to a separate XLSX sheet; "
            "repeat the option or separate MACs with commas"
        ),
    )
    parser.add_argument(
        "--xlsx",
        type=Path,
        metavar="PATH",
        help="also save the report as an XLSX workbook",
    )
    args = parser.parse_args()
    if args.start_time and args.end_time and args.start_time > args.end_time:
        parser.error("--start-time must be earlier than or equal to --end-time")
    if args.top_error_clients is not None and args.top_error_clients < 1:
        parser.error("--top-error-clients must be greater than 0")
    if args.top_error_client_logs is not None:
        if args.top_error_client_logs < 1:
            parser.error("--top-error-client-logs must be greater than 0")
        if args.top_error_clients is None:
            parser.error("--top-error-client-logs requires --top-error-clients")
        if args.top_error_client_logs > args.top_error_clients:
            parser.error(
                "--top-error-client-logs must be less than or equal to "
                "--top-error-clients"
            )
        if args.xlsx is None:
            parser.error("--top-error-client-logs requires --xlsx")
    if args.client_logs:
        args.client_logs = tuple(
            dict.fromkeys(
                client
                for client_group in args.client_logs
                for client in client_group
            )
        )
        if args.xlsx is None:
            parser.error("--client-logs requires --xlsx")
    return args


def main() -> int:
    """Run the command-line report."""
    args = parse_args()
    try:
        with args.log_file.open(encoding="utf-8", errors="replace") as stream:
            by_mac, by_name = extract_bssids(stream, args.ssid)
        mapped, blank, multiple, changed, backup_path = update_inventory(
            args.inventory,
            args.inventory,
            f"{args.ssid} BSSID",
            by_mac,
            by_name,
        )
        print(f"SSID BSSID mapping: {mapped} APs mapped, {blank} blank, "
              f"{multiple} with multiple BSSIDs, {changed} changed")
        if backup_path:
            print(f"Inventory backup saved to: {backup_path}")
        with args.inventory.open(encoding="utf-8-sig", newline="") as stream:
            inventory = load_inventory(stream)
        with args.log_file.open(encoding="utf-8", errors="replace") as stream:
            summary = summarize(
                stream,
                args.ssid,
                args.start_time,
                args.end_time,
                args.client_logs,
                inventory,
            )
    except (OSError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    if not summary.counts and args.client_logs is None:
        print("No onboarding attempts found.", file=sys.stderr)
        return 2

    print(f"SSID: {args.ssid}\n")
    if summary.start_time and summary.stop_time:
        time_format = "%Y-%m-%d %H:%M:%S %Z"
        print(f"Start time: {summary.start_time.strftime(time_format)}")
        print(f"Stop time:  {summary.stop_time.strftime(time_format)}\n")
    print("AAA authentication\n")
    print(format_aaa_summary(summary.aaa_transactions))
    if args.client_logs:
        print("\nAAA by requested client\n")
        print(format_aaa_clients(summary.aaa_transactions, args.client_logs))
    platform_counts = aggregate_by_platform(
        summary.counts, summary.association_events, inventory
    )
    if args.top_error_client_logs is not None:
        ranked_client_logs = top_error_client_macs(
            platform_counts, args.top_error_client_logs
        )
        client_logs = tuple(
            dict.fromkeys((*(args.client_logs or ()), *ranked_client_logs))
        )
        if client_logs != (args.client_logs or ()):
            try:
                with args.log_file.open(encoding="utf-8", errors="replace") as stream:
                    summary = summarize(
                        stream,
                        args.ssid,
                        args.start_time,
                        args.end_time,
                        client_logs,
                        inventory,
                    )
            except OSError as error:
                print(f"error: {error}", file=sys.stderr)
                return 1
        args.client_logs = client_logs
    if summary.counts:
        print("Per AP\n")
        print(format_ap_table(summary.counts, inventory))
        print("\nBy platform\n")
        print(format_platform_table(platform_counts))
        print("\nFailure reasons by platform\n")
        print(format_failure_table(platform_counts, args.show_macs))
        if args.top_error_clients is not None:
            print(f"\nTop {args.top_error_clients} failing clients by platform\n")
            print(format_top_error_clients(platform_counts, args.top_error_clients))
    else:
        print("No onboarding attempts found for the selected SSID and time range.")
    if args.xlsx:
        try:
            save_xlsx(
                args.xlsx,
                summary,
                inventory,
                platform_counts,
                args.ssid,
                args.show_macs,
                args.top_error_clients,
                args.client_logs,
            )
        except (OSError, ImportError, RuntimeError) as error:
            print(f"error: could not save XLSX: {error}", file=sys.stderr)
            return 1
        print(f"\nXLSX saved to: {args.xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())